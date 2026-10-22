#!/usr/bin/env python3
"""Tests for pbudget: `python -m unittest -v` (or `python test_pbudget.py`).

Roll-ups, interpolation and margins are checked against values worked out independently of
the code under test. Only the standard library, NumPy and PyYAML are needed.
"""
import contextlib
import io
import json
import pathlib
import re
import tempfile
import unittest

import numpy as np
import yaml

import yaml2budget as pb


def write(directory, name, text):
    path = pathlib.Path(directory) / name
    path.write_text(text)
    return path


def build(text, directory=None, name='b.yaml', rollup='rss', files=()):
    """Evaluate a budget given as YAML text. Returns (budget, grid, warnings)."""
    directory = directory or tempfile.mkdtemp()
    for filename, content in files:
        write(directory, filename, content)
    path = write(directory, name, text)
    warnings = io.StringIO()
    with contextlib.redirect_stderr(warnings):
        budget, grid = pb.load_budget(path, rollup)
    return budget, grid, warnings.getvalue()


def node(budget, name):
    return next(n for n in pb.walk(budget) if n['name'] == name)


TREE = """
name: T
rollup: {rollup}
units: u
children:
  - name: A
    alloc: {alloc_a}
    children:
      - {{name: A1, cbe: 3}}
      - {{name: A2, cbe: 4}}
  - name: B
    children:
      - {{name: B1, cbe: 2}}
      - name: B2
        children:
          - {{name: B2a, cbe: 5}}
          - {{name: B2b, cbe: 12}}
"""


class Rollups(unittest.TestCase):
    """Scalars combined by each operator, at every level of the hierarchy."""

    def rolled(self, rollup, alloc_a='6'):
        budget, _, _ = build(TREE.format(rollup=rollup, alloc_a=alloc_a))
        return budget

    def test_rss(self):
        b = self.rolled('rss')
        self.assertAlmostEqual(node(b, 'A')['cbe'], 5.0)             # sqrt(3^2 + 4^2)
        self.assertAlmostEqual(node(b, 'B2')['cbe'], 13.0)           # sqrt(5^2 + 12^2)
        self.assertAlmostEqual(node(b, 'B')['cbe'], np.hypot(2, 13))
        self.assertAlmostEqual(b['cbe'], np.sqrt(5 ** 2 + 2 ** 2 + 13 ** 2))

    def test_sum(self):
        b = self.rolled('sum')
        self.assertAlmostEqual(node(b, 'A')['cbe'], 7)
        self.assertAlmostEqual(node(b, 'B2')['cbe'], 17)
        self.assertAlmostEqual(node(b, 'B')['cbe'], 19)
        self.assertAlmostEqual(b['cbe'], 26)

    def test_product(self):
        b = self.rolled('product')
        self.assertAlmostEqual(node(b, 'A')['cbe'], 12)
        self.assertAlmostEqual(node(b, 'B2')['cbe'], 60)
        self.assertAlmostEqual(node(b, 'B')['cbe'], 120)
        self.assertAlmostEqual(b['cbe'], 1440)

    def test_child_override_is_inherited_by_its_own_subtree(self):
        budget, _, _ = build("""
name: T
rollup: rss
children:
  - name: A
    rollup: sum
    children:
      - {name: A1, cbe: 3}
      - name: A2
        children:                       # inherits sum from A, not rss from the root
          - {name: A2a, cbe: 1}
          - {name: A2b, cbe: 2}
  - {name: B, cbe: 4}
""")
        self.assertAlmostEqual(node(budget, 'A2')['cbe'], 3)
        self.assertAlmostEqual(node(budget, 'A')['cbe'], 6)
        self.assertAlmostEqual(budget['cbe'], np.hypot(6, 4))

    def test_default_rollup_applies_only_when_the_file_is_silent(self):
        silent = 'name: T\nchildren: [{name: A, cbe: 3}, {name: B, cbe: 4}]\n'
        self.assertAlmostEqual(build(silent, rollup='sum')[0]['cbe'], 7)
        self.assertAlmostEqual(build(silent, rollup='rss')[0]['cbe'], 5)
        stated = 'name: T\nrollup: product\nchildren: [{name: A, cbe: 3}, {name: B, cbe: 4}]\n'
        self.assertAlmostEqual(build(stated, rollup='sum')[0]['cbe'], 12)

    def test_single_child_parent_equals_its_child(self):
        budget, _, _ = build('name: T\nrollup: rss\nchildren: [{name: A, cbe: 7}]\n')
        self.assertAlmostEqual(budget['cbe'], 7)


class Counts(unittest.TestCase):
    """`count: N` enters a term N times into its parent's roll-up; its own values stay per copy."""

    def rolled(self, rollup, count=3, extra=''):
        budget, _, _ = build(f"""
name: T
rollup: {rollup}
children:
  - {{name: A, cbe: 2, alloc: 2.5, count: {count}{extra}}}
  - {{name: B, cbe: 5}}
""")
        return budget

    def test_each_operator_sees_n_identical_terms(self):
        expected = {'product': 2 * 2 * 2 * 5, 'sum': 2 + 2 + 2 + 5,
                    'rss': np.sqrt(3 * 2 ** 2 + 5 ** 2),
                    'inverse_rss': 1 / np.sqrt(3 / 2 ** 2 + 1 / 5 ** 2)}
        for rollup, value in expected.items():
            with self.subTest(rollup):
                self.assertAlmostEqual(self.rolled(rollup)['cbe'], value)

    def test_repeats_match_the_rollup_of_copies(self):
        v = np.array([0.5, 0.9, 2.0])
        for rollup, combine in pb.ROLLUPS.items():
            with self.subTest(rollup):
                np.testing.assert_allclose(pb.REPEATS[rollup](v, 4), combine([v] * 4))

    def test_the_node_keeps_its_per_copy_values(self):
        a = node(self.rolled('product'), 'A')
        self.assertEqual((a['cbe'], a['alloc']), (2, 2.5))

    def test_a_vector_term_is_raised_to_the_power(self):
        budget, _, _ = build("""
name: T
rollup: product
wavelength: [400, 800, 400]
children: [{name: Mirrors, cbe: [[400, 0.9], [800, 0.8]], count: 5}]
""")
        np.testing.assert_allclose(budget['cbe'], [0.9 ** 5, 0.8 ** 5])

    def test_a_repeated_subtree_counts_as_a_whole(self):
        budget, _, _ = build("""
name: T
rollup: sum
children:
  - name: Camera
    count: 2
    children: [{name: Lens, cbe: 3}, {name: Body, cbe: 4}]
""")
        self.assertAlmostEqual(node(budget, 'Camera')['cbe'], 7)
        self.assertAlmostEqual(budget['cbe'], 14)

    def test_aliases(self):
        for key in ('multiplier', 'quantity'):
            budget, _, _ = build(f'name: T\nrollup: sum\n'
                                 f'children: [{{name: A, cbe: 2, {key}: 4}}]\n')
            self.assertAlmostEqual(budget['cbe'], 8)

    def test_invalid_counts_are_errors(self):
        for count in ('0', '-2', 'many', 'true'):
            with self.subTest(count), self.assertRaises(ValueError):
                self.rolled('sum', count)

    def test_the_root_cannot_be_repeated(self):
        with self.assertRaises(ValueError):
            build('name: T\ncount: 2\nrollup: sum\nchildren: [{name: A, cbe: 1}]\n')

    def test_the_page_gets_the_count_and_not_a_column(self):
        page = pb.flatten(self.rolled('sum'), None)
        self.assertEqual([n['count'] for n in page['nodes']], [1, 3, 1])
        self.assertEqual(page['columns'], [])

    def test_derived_budgets_stay_consistent(self):
        directory = tempfile.mkdtemp()
        budget, grid, _ = build("""
name: Blur
rollup: rss
wavelength: [400, 800, 400]
children: [{name: A, cbe: 2, count: 4}, {name: B, cbe: 3}]
derived: [{name: R, scale: 1000, invert: true, rollup: inverse_rss, better: higher}]
""", directory)
        derived = pb.derive(budget, budget['derived'][0], grid, pathlib.Path(directory))
        self.assertAlmostEqual(float(np.ravel(budget['cbe'])[0]), 5)           # sqrt(4*4 + 9)
        self.assertAlmostEqual(float(np.ravel(derived['cbe'])[0]), 1000 / 5)   # scale / blur
        recomputed = pb.ROLLUPS['inverse_rss'](
            [pb.REPEATS['inverse_rss'](node(derived, 'A')['cbe'], 4), node(derived, 'B')['cbe']])
        self.assertAlmostEqual(float(np.ravel(recomputed)[0]), 200)


class VectorRollups(unittest.TestCase):
    """The same operators with wavelength-dependent terms, checked per grid point."""

    CSV = 'wavelength_nm,value\n400,1\n600,3\n800,5\n'

    def test_vectors_and_scalars_mix(self):
        budget, grid, _ = build("""
name: T
rollup: rss
wavelength: [400, 800, 100]
children:
  - {name: A, cbe: v.csv}
  - {name: B, cbe: 4}
  - {name: C, cbe: [[400, 0], [800, 4]]}
""", files=[('v.csv', self.CSV)])
        np.testing.assert_allclose(grid, [400, 500, 600, 700, 800])
        a = node(budget, 'A')['cbe']
        c = node(budget, 'C')['cbe']
        np.testing.assert_allclose(a, [1, 2, 3, 4, 5])               # linear interpolation
        np.testing.assert_allclose(c, [0, 1, 2, 3, 4])
        np.testing.assert_allclose(budget['cbe'], np.sqrt(a ** 2 + 16 + c ** 2))

    def test_product_of_vectors(self):
        budget, _, _ = build("""
name: T
rollup: product
wavelength: [400, 800, 200]
children:
  - {name: A, cbe: v.csv}
  - {name: B, cbe: 0.5}
""", files=[('v.csv', self.CSV)])
        np.testing.assert_allclose(budget['cbe'], [0.5, 1.5, 2.5])

    def test_scalar_stays_scalar_when_no_table_is_present(self):
        budget, grid, _ = build('name: T\nrollup: sum\nchildren: [{name: A, cbe: 2}]\n')
        self.assertIsNone(grid)
        self.assertEqual(np.ndim(budget['cbe']), 0)

    def test_values_outside_a_table_are_clamped_with_a_warning(self):
        budget, grid, warnings = build("""
name: T
rollup: sum
wavelength: [300, 900, 300]
children: [{name: A, cbe: v.csv}]
""", files=[('v.csv', self.CSV)])
        np.testing.assert_allclose(budget['cbe'], [1, 3, 5])         # held flat past both ends
        self.assertIn('clamped to grid', warnings)


class Grids(unittest.TestCase):
    def test_explicit_grid_includes_its_endpoint(self):
        _, grid, _ = build('name: T\nwavelength: [400, 1000, 10]\nchildren: [{name: A, cbe: 1}]\n')
        self.assertEqual((grid[0], grid[-1], len(grid)), (400, 1000, 61))

    def test_inferred_grid_covers_the_common_range_with_a_1_2_5_step(self):
        files = [('a.csv', 'w,v\n380,1\n1050,1\n'), ('b.csv', 'w,v\n400,2\n1000,2\n')]
        _, grid, _ = build("""
name: T
rollup: sum
children: [{name: A, cbe: a.csv}, {name: B, cbe: b.csv}]
""", files=files)
        step = np.diff(grid)
        self.assertTrue(np.allclose(step, step[0]))
        self.assertIn(step[0], (1, 2, 5, 10, 20, 50))
        self.assertGreaterEqual(len(grid), pb.MIN_POINTS)
        self.assertGreaterEqual(grid[0], 400)                        # inside both tables
        self.assertLessEqual(grid[-1], 1000)

    def test_tables_that_do_not_overlap_are_an_error(self):
        files = [('a.csv', 'w,v\n300,1\n400,1\n'), ('b.csv', 'w,v\n900,2\n1000,2\n')]
        with self.assertRaises(ValueError):
            build('name: T\nrollup: sum\n'
                  'children: [{name: A, cbe: a.csv}, {name: B, cbe: b.csv}]\n', files=files)


class Margins(unittest.TestCase):
    def test_lower_is_better(self):
        difference, fraction = pb.margins(cbe=8.0, alloc=10.0, better='lower')
        self.assertAlmostEqual(difference, 2.0)
        self.assertAlmostEqual(fraction, 0.2)

    def test_higher_is_better(self):
        difference, fraction = pb.margins(cbe=0.6, alloc=0.5, better='higher')
        self.assertAlmostEqual(difference, 0.1)
        self.assertAlmostEqual(fraction, 0.2)

    def test_exceeded_allocation_is_negative_both_ways(self):
        for better, cbe, alloc in (('lower', 12.0, 10.0), ('higher', 0.4, 0.5)):
            difference, fraction = pb.margins(cbe, alloc, better)
            self.assertLess(difference, 0)
            self.assertLess(fraction, 0)

    def test_vector_margin_is_elementwise(self):
        difference, fraction = pb.margins(np.array([2.0, 5.0]), np.array([4.0, 4.0]), 'lower')
        np.testing.assert_allclose(difference, [2.0, -1.0])
        np.testing.assert_allclose(fraction, [0.5, -0.25])

    def test_no_allocation_gives_no_margin(self):
        self.assertEqual(pb.margins(1.0, None, 'lower'), (None, None))

    def test_better_defaults_to_the_direction_of_the_rollup(self):
        product, _, _ = build('name: T\nrollup: product\nchildren: [{name: A, cbe: 1}]\n')
        summed, _, _ = build('name: T\nrollup: sum\nchildren: [{name: A, cbe: 1}]\n')
        self.assertEqual(product['better'], 'higher')
        self.assertEqual(summed['better'], 'lower')

    def test_page_margin_matches_the_python_one(self):
        budget, _, _ = build("""
name: T
rollup: sum
alloc: 10
children: [{name: A, cbe: 3, alloc: 5}, {name: B, cbe: 4}]
""")
        page = pb.flatten(budget, None)
        root = page['nodes'][0]
        difference, fraction = pb.margins(root['cbe'], root['alloc'], page['better'])
        self.assertAlmostEqual(root['cbe'], 7)
        self.assertAlmostEqual(difference, 3)
        self.assertAlmostEqual(fraction, 0.3)


class SuppliedValuesAreReplaced(unittest.TestCase):
    YAML = """
name: T
rollup: sum
alloc: 10
cbe: 99
margin: 42
margin_frac: 0.99
children: [{name: A, cbe: 3}, {name: B, cbe: 4}]
"""

    def setUp(self):
        self.budget, _, self.warnings = build(self.YAML)

    def test_parent_cbe_is_recomputed(self):
        self.assertAlmostEqual(self.budget['cbe'], 7)
        self.assertIn('supplied cbe 99 replaced by computed 7', self.warnings)

    def test_margins_are_dropped_and_warned_about(self):
        self.assertNotIn('margin', self.budget)
        self.assertNotIn('margin_frac', self.budget)
        self.assertIn('supplied margin 42 replaced by computed 3', self.warnings)
        self.assertIn('supplied margin_frac 0.99 replaced by computed 0.3', self.warnings)

    def test_replaced_keys_do_not_become_columns(self):
        self.assertEqual(pb.flatten(self.budget, None)['columns'], [])


class Schema(unittest.TestCase):
    def test_named_sub_mappings_become_children(self):
        budget = pb.normalize(yaml.safe_load("""
telescope:
  optics: {cbe: 1}
  mount: {cbe: 2}
"""), 'file-stem')
        self.assertEqual(budget['name'], 'file-stem')
        self.assertEqual([c['name'] for c in budget['children']], ['telescope'])
        self.assertEqual([c['name'] for c in budget['children'][0]['children']],
                         ['optics', 'mount'])

    def test_aliases(self):
        budget = pb.normalize(yaml.safe_load(
            'name: T\nallocation: 5\ndescription: hi\nchildren: [{name: A, spec: 2, cbe: 1}]\n'))
        self.assertEqual(budget['alloc'], 5)
        self.assertEqual(budget['notes'], 'hi')
        self.assertEqual(budget['children'][0]['alloc'], 2)

    def test_list_and_mapping_children_can_be_mixed(self):
        budget = pb.normalize(yaml.safe_load(
            'name: T\nchildren: [{name: A, cbe: 1}]\nextra_branch: {cbe: 2}\n'))
        self.assertEqual([c['name'] for c in budget['children']], ['A', 'extra_branch'])

    def test_wrapper_root_is_merged_into_its_only_child(self):
        budget, _, _ = build("""
name: Contrast
description: about this budget
total:
  rollup: sum
  alloc: 10
  a: {cbe: 1}
  b: {cbe: 2}
""")
        self.assertEqual(budget['name'], 'Contrast')                 # keeps the file's title
        self.assertEqual(budget['description'], 'about this budget')
        self.assertEqual(budget['alloc'], 10)                        # and the wrapper's values
        self.assertAlmostEqual(budget['cbe'], 3)
        self.assertEqual([c['name'] for c in budget['children']], ['a', 'b'])

    def test_leaf_without_a_cbe_is_an_error(self):
        with self.assertRaises(ValueError):
            build('name: T\nchildren: [{name: A, notes: todo}]\n')


class Derived(unittest.TestCase):
    SPATIAL = """
name: Spatial
rollup: rss
units: urad
wavelength: [400, 800, 200]
alloc: 20
children:
  - {name: A, cbe: 3, alloc: 6}
  - {name: B, cbe: 4}
derived:
  - {name: Spectral, units: nm, scale: [[400, 0.1], [800, 0.3]], alloc: 1.0}
"""

    def setUp(self):
        directory = tempfile.mkdtemp()
        self.budget, self.grid, _ = build(self.SPATIAL, directory)
        spec = self.budget['derived'][0]
        self.derived = pb.derive(self.budget, spec, self.grid, pathlib.Path(directory))

    def test_every_value_is_scaled(self):
        scale = np.array([0.1, 0.2, 0.3])
        np.testing.assert_allclose(self.derived['cbe'], 5 * scale)   # rss(3, 4) = 5
        np.testing.assert_allclose(node(self.derived, 'A')['cbe'], 3 * scale)
        np.testing.assert_allclose(node(self.derived, 'A')['alloc'], 6 * scale)

    def test_root_allocation_can_be_overridden_and_metadata_follows(self):
        self.assertEqual(self.derived['alloc'], 1.0)
        self.assertEqual(self.derived['name'], 'Spectral')
        self.assertEqual(self.derived['units'], 'nm')

    def test_scaling_commutes_with_the_rollup(self):
        scaled_children = [node(self.derived, name)['cbe'] for name in ('A', 'B')]
        np.testing.assert_allclose(self.derived['cbe'], np.hypot(*scaled_children))

    def test_the_source_budget_is_untouched(self):
        self.assertAlmostEqual(float(np.ravel(self.budget['cbe'])[0]), 5)


class ReciprocalBudgets(unittest.TestCase):
    """Resolving power: R = lambda/dlambda, an inverted resolution budget."""

    SPATIAL = """
name: Blur
rollup: rss
units: arcsec
wavelength: [400, 800, 200]
alloc: 3.0
children:
  - {name: Optics, cbe: 2.0, alloc: 2.5}
  - {name: Jitter, cbe: 1.0}
derived:
  - name: R
    units: ''
    scale: [[400, 4000], [800, 8000]]
    invert: true
    rollup: inverse_rss
    better: higher
    alloc: 1500
"""

    def setUp(self):
        directory = tempfile.mkdtemp()
        self.budget, self.grid, _ = build(self.SPATIAL, directory)
        self.derived = pb.derive(self.budget, self.budget['derived'][0], self.grid,
                                 pathlib.Path(directory))

    def test_operator_combines_reciprocals_in_quadrature(self):
        value = pb.ROLLUPS['inverse_rss']([np.array([2000.0]), np.array([3000.0])])
        np.testing.assert_allclose(value, 1 / np.sqrt(1 / 2000 ** 2 + 1 / 3000 ** 2))

    def test_it_defaults_to_higher_is_better(self):
        budget, _, _ = build('name: T\nrollup: inverse_rss\nchildren: [{name: A, cbe: 100}]\n')
        self.assertEqual(budget['better'], 'higher')

    def test_each_value_becomes_scale_over_value(self):
        scale = np.array([4000.0, 6000.0, 8000.0])
        np.testing.assert_allclose(node(self.derived, 'Optics')['cbe'], scale / 2)
        np.testing.assert_allclose(node(self.derived, 'Optics')['alloc'], scale / 2.5)

    def test_the_root_is_the_inverse_rss_of_the_inverted_leaves(self):
        leaves = [node(self.derived, name)['cbe'] for name in ('Optics', 'Jitter')]
        np.testing.assert_allclose(self.derived['cbe'],
                                   1 / np.sqrt(sum(1 / v ** 2 for v in leaves)))
        np.testing.assert_allclose(self.derived['cbe'],                  # == scale / rss(blur)
                                   np.array([4000.0, 6000.0, 8000.0]) / np.hypot(2.0, 1.0))

    def test_the_page_reports_it_as_a_reciprocal_budget(self):
        page = pb.flatten(self.derived, self.grid)
        self.assertEqual((page['better'], page['units']), ('higher', ''))
        self.assertEqual([n['op'] for n in page['nodes']], ['inverse_rss', None, None])
        self.assertEqual(page['nodes'][0]['alloc'], 1500)
        self.assertEqual(page['columns'], [])                            # `invert` is not a column

    def test_margins_run_the_higher_is_better_way(self):
        difference, fraction = pb.margins(2400.0, 2000.0, 'higher')
        self.assertAlmostEqual(difference, 400.0)
        self.assertAlmostEqual(fraction, 0.2)


class PageData(unittest.TestCase):
    def setUp(self):
        self.budget, self.grid, _ = build("""
name: T
rollup: sum
levels: [System, Part]
units: kg
children:
  - name: A
    source: vendor
    children: [{name: A1, cbe: 1, tbr: true}]
  - {name: B, cbe: 2, notes: n}
""")
        self.page = pb.flatten(self.budget, self.grid)

    def test_nodes_are_depth_first_with_parent_links(self):
        self.assertEqual([(n['name'], n['depth']) for n in self.page['nodes']],
                         [('T', 0), ('A', 1), ('A1', 2), ('B', 1)])
        self.assertEqual([n['parent'] for n in self.page['nodes']], [None, 0, 1, 0])

    def test_parents_carry_their_operator_and_leaves_do_not(self):
        self.assertEqual([n['op'] for n in self.page['nodes']], ['sum', 'sum', None, None])

    def test_extra_keys_become_columns_in_first_seen_order(self):
        self.assertEqual(self.page['columns'], ['source', 'tbr'])
        self.assertEqual(self.page['nodes'][1]['extra'], {'source': 'vendor'})
        self.assertEqual(self.page['nodes'][2]['extra'], {'tbr': True})

    def test_settings_reach_the_page(self):
        self.assertEqual((self.page['units'], self.page['levels']), ('kg', ['System', 'Part']))
        self.assertEqual((self.page['better'], self.page['warn']), ('lower', 0.1))

    def test_values_are_json_safe(self):
        json.dumps(self.page)                                        # raises on numpy types
        self.assertIsInstance(self.page['nodes'][0]['cbe'], float)

    def test_rounding(self):
        self.assertEqual(pb.rounded(1 / 3), 0.33333)
        self.assertEqual(pb.rounded(np.array([1 / 3, 2 / 3])), [0.33333, 0.66667])
        self.assertIsNone(pb.rounded(None))


class Views(unittest.TestCase):
    def views(self, line=''):
        budget, grid, _ = build(f'name: T\nrollup: sum\n{line}children: [{{name: A, cbe: 1}}]\n')
        return pb.flatten(budget, grid)['views']

    def test_all_three_by_default(self):
        self.assertEqual(self.views(), ['table', 'flow', 'plot'])

    def test_a_list_selects_and_orders_them(self):
        self.assertEqual(self.views('views: [flowchart, table]\n'), ['flow', 'table'])

    def test_a_single_name_is_allowed(self):
        self.assertEqual(self.views('views: table\n'), ['table'])

    def test_a_map_of_flags_selects_the_true_ones(self):
        self.assertEqual(self.views('views: {table: true, flowchart: false, spectrum: true}\n'),
                         ['table', 'plot'])

    def test_page_ids_are_accepted_too(self):
        self.assertEqual(self.views('views: [flow, plot]\n'), ['flow', 'plot'])

    def test_unknown_view_is_an_error(self):
        with self.assertRaises(ValueError):
            self.views('views: [table, histogram]\n')

    def test_disabling_everything_is_an_error(self):
        with self.assertRaises(ValueError):
            self.views('views: []\n')

    def test_views_does_not_become_a_column(self):
        budget, grid, _ = build('name: T\nviews: [table]\nchildren: [{name: A, cbe: 1}]\n')
        self.assertEqual(pb.flatten(budget, grid)['columns'], [])

    def test_derived_budgets_can_override(self):
        directory = tempfile.mkdtemp()
        budget, grid, _ = build("""
name: Spatial
rollup: rss
wavelength: [400, 800, 200]
views: [table, spectrum]
children: [{name: A, cbe: 3}]
derived:
  - {name: Spectral, scale: 2, views: [table]}
""", directory)
        derived = pb.derive(budget, budget['derived'][0], grid, pathlib.Path(directory))
        self.assertEqual(pb.flatten(budget, grid)['views'], ['table', 'plot'])
        self.assertEqual(pb.flatten(derived, grid)['views'], ['table'])


class Page(unittest.TestCase):
    """The generated file: one entry per budget, data readable, libraries wired up."""

    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.mkdtemp()
        write(cls.directory, 'w.csv', 'w,v\n400,0.8\n800,0.4\n')
        write(cls.directory, 'a.yaml', """
name: Throughput
rollup: product
wavelength: [400, 800, 200]
alloc: 0.3
children: [{name: A, cbe: w.csv}, {name: B, cbe: 0.9}]
""")
        write(cls.directory, 'b.yaml', 'name: Mass\nrollup: sum\nchildren: [{name: A, cbe: 2}]\n')
        with contextlib.redirect_stderr(io.StringIO()):
            cls.pages = pb.load_all(cls.directory)
            cls.html = pb.render(cls.pages)

    def embedded(self):
        match = re.search(r'const BUDGETS = (\[.*?\]);\n', self.html, re.S)
        return json.loads(match.group(1))

    def test_one_page_per_file_in_filename_order(self):
        self.assertEqual([(p['file'], p['name']) for p in self.pages],
                         [('a.yaml', 'Throughput'), ('b.yaml', 'Mass')])

    def test_embedded_data_matches_an_independent_calculation(self):
        throughput = next(p for p in self.embedded() if p['name'] == 'Throughput')
        expected = np.interp([400, 600, 800], [400, 800], [0.8, 0.4]) * 0.9
        np.testing.assert_allclose(throughput['nodes'][0]['cbe'], expected, rtol=1e-4)
        self.assertEqual(throughput['grid'], [400, 600, 800])
        self.assertEqual(throughput['better'], 'higher')             # product -> higher is better

    def test_scalar_budget_has_no_grid(self):
        self.assertIsNone(next(p for p in self.embedded() if p['name'] == 'Mass')['grid'])

    def test_libraries_are_referenced_once_each(self):
        for url in pb.CDN:
            self.assertEqual(self.html.count(url), 1)

    def test_inlined_libraries_replace_the_cdn_tags(self):
        library = pathlib.Path(write(self.directory, 'lib.js', 'window.stub = 1;\n'))
        html = pb.render(self.pages, inline=[library])
        self.assertIn('window.stub = 1;', html)
        self.assertNotIn(pb.CDN[0], html)

    def test_closing_tags_in_the_data_cannot_end_the_script(self):
        page = dict(self.pages[0], description='</script><script>alert(1)</script>')
        self.assertNotIn('</script><script>alert(1)', pb.render([page]))


def _optional(name):
    try:
        __import__(name)
        return False
    except ImportError:
        return True


@unittest.skipIf(_optional('playwright') or _optional('openpyxl'),
                 'needs playwright and openpyxl (pip install playwright openpyxl)')
class ExcelExport(unittest.TestCase):
    """The page's Excel button, driven in a headless browser and read back with openpyxl."""

    @classmethod
    def setUpClass(cls):
        import openpyxl
        from playwright.sync_api import sync_playwright
        directory = pathlib.Path(tempfile.mkdtemp())
        write(directory, 'b.yaml', """
name: Demo
rollup: sum
units: kg
levels: [System, Subsystem, Part]
wavelength: [400, 600, 100]
alloc: 10
children:
  - name: Optics
    alloc: 6
    children: [{name: Lens, cbe: [[400, 1], [600, 3]], notes: n}, {name: Mount, cbe: 2}]
  - {name: Electronics, cbe: 4, alloc: 3}
""")
        page = directory / 'page.html'
        with contextlib.redirect_stderr(io.StringIO()):
            page.write_text(pb.render(pb.load_all(directory)))
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            context = browser.new_context(accept_downloads=True)
            tab = context.new_page()
            cls.errors = []
            tab.on('pageerror', lambda e: cls.errors.append(str(e)))
            tab.goto(page.as_uri())
            tab.wait_for_timeout(200)
            with tab.expect_download(timeout=30000) as download:
                tab.click('#exXlsx')
            download.value.save_as(directory / 'out.xlsx')
            browser.close()
        cls.book = openpyxl.load_workbook(directory / 'out.xlsx')
        cls.sheet = cls.book['Demo']

    def test_no_page_errors(self):
        self.assertEqual(self.errors, [])

    def test_one_sheet_per_view_of_the_data(self):
        self.assertEqual(self.book.sheetnames, ['Demo', 'Spectrum'])

    def rows(self):
        """{entry name: its row}, the name being whichever level column holds it."""
        return {next(c.value for c in row[:3] if c.value): row
                for row in self.sheet.iter_rows(min_row=3)}

    def test_header_is_frozen_below_the_title(self):
        self.assertEqual(self.sheet.freeze_panes, 'A3')
        self.assertEqual([c.value for c in self.sheet[2]][:5],
                         ['System', 'Subsystem', 'Part', 'Roll-up', 'CBE [kg]'])

    def test_values_are_numbers_at_the_current_wavelength(self):
        rows = self.rows()
        self.assertAlmostEqual(rows['Lens'][4].value, 1)          # λ starts at 400 nm
        self.assertAlmostEqual(rows['Demo'][4].value, 7)          # 1 + 2 + 4
        self.assertAlmostEqual(rows['Optics'][4].value, 3)

    def test_rollup_rows_are_bold_shaded_and_grouped(self):
        root, optics, lens = (self.sheet.cell(row=r, column=c) for r, c in ((3, 1), (4, 2), (5, 3)))
        self.assertTrue(root.font.b and optics.font.b)
        self.assertFalse(lens.font.b)
        self.assertEqual(optics.fill.fgColor.rgb, 'FFF0EFEC')
        self.assertEqual([self.sheet.row_dimensions[r].outline_level for r in (3, 4, 5)], [0, 1, 2])

    def test_margins_are_percentages_coloured_by_sign(self):
        margin = {name: row[6] for name, row in self.rows().items()}
        self.assertAlmostEqual(margin['Demo'].value, 0.3)         # 1 - 7/10
        self.assertTrue(margin['Demo'].number_format.startswith('+0.0%'))
        self.assertEqual(margin['Demo'].font.color.rgb, 'FF0A7A0A')
        self.assertAlmostEqual(margin['Electronics'].value, -1 / 3)
        self.assertEqual(margin['Electronics'].font.color.rgb, 'FFC0392B')

    def test_spectrum_sheet_holds_the_whole_band(self):
        spectrum = self.book['Spectrum']
        self.assertEqual([c.value for c in spectrum[2]][:2], ['Wavelength [nm]', 'Demo'])
        self.assertEqual([r[0].value for r in spectrum.iter_rows(min_row=3)], [400, 500, 600])
        self.assertAlmostEqual(spectrum.cell(row=5, column=2).value, 9)   # 3 + 2 + 4 at 600 nm


if __name__ == '__main__':
    unittest.main(verbosity=2)
